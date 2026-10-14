"""KPI Report: one employee's KPI as an Excel sheet to print, sign and file.

Laid out like the monthly OKR sheet HR uses (docs/Vairable _ July 2026.xlsx): the company
header, who the employee is and who they report to, the score and how it is worked out,
the projects (objective / expected outcome) and tasks behind it, remarks from the line
manager, and lines for the employee's and the reporting manager's signatures.

Nothing is stored: the figures are the live KPI of routers/kpi.py at the moment of download.
"""
import re
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import models

LOGO = Path(__file__).parent / "assets" / "anwars-logo.jpg"
COMPANY = "ANWAR GROUP OF INDUSTRIES"
LAST = 15  # columns A..O

NAVY, LIGHT, PALE = "1F3864", "D9E2F3", "F2F2F2"
RATING_TEXT = {"excellent": "Excellent", "good": "Good", "fair": "Fair",
               "needs_attention": "Needs attention", "no_tasks": "Nothing to score yet"}
RATING_FILL = {"excellent": "C6EFCE", "good": "DDEBF7", "fair": "FFEB9C", "needs_attention": "FFC7CE"}
ROLE_TEXT = {"responsible": "Responsible", "accountable": "Accountable", "reviewer": "Reviewer",
             "project_manager": "Project Manager", "tasks": "All tasks"}
STATE_TEXT = {"completed": "Completed", "overdue": "Past due", "running": "Running", "not_due": "Not counted yet"}

_thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _plain(text: Optional[str]) -> str:
    """The light markdown of the rich-text fields (components/RichText.tsx) as plain text."""
    if not text:
        return ""
    text = re.sub(r"\[([^\]\n]+)\]\(([^)\s]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("++", "")
    text = re.sub(r"(?<![\w*])\*(?!\s)([^*\n]+?)\*(?![\w*])", r"\1", text)
    return text.strip()


def _label(value: Optional[str]) -> str:
    return (value or "").replace("_", " ").title() if value else "—"


def _num(value) -> str:
    return "—" if value is None else f"{value:g}"


def _pct(value) -> str:
    return "—" if value is None else f"{value:g}%"


def _day(d: Optional[date]) -> str:
    return d.strftime("%d %b %Y") if d else "—"


def _on_time(row: dict) -> str:
    """A task's on-time completion, in words."""
    if row["state"] == "completed":
        return "On time" if row["on_time"] is not False else "Late"
    return {"overdue": "Overdue", "running": "In progress", "not_due": "Not due yet"}[row["state"]]


class _Sheet:
    def __init__(self, ws):
        self.ws, self.row = ws, 1

    def cell(self, col: int, value, *, bold=False, fill=None, align=WRAP, size=10, color=None, border=True,
             to: Optional[int] = None, number_format: Optional[str] = None):
        c = self.ws.cell(self.row, col, value)
        c.font = Font(bold=bold, size=size, color=color)
        c.alignment = align
        if number_format:
            c.number_format = number_format
        last = to or col
        if last > col:
            self.ws.merge_cells(start_row=self.row, start_column=col, end_row=self.row, end_column=last)
        for i in range(col, last + 1):
            cc = self.ws.cell(self.row, i)
            if fill:
                cc.fill = PatternFill("solid", fgColor=fill)
            if border:
                cc.border = BOX
        return c

    def height(self, *texts_and_widths, base=15):
        """Row tall enough for the longest wrapped text (Excel does not grow merged rows itself)."""
        lines = 1
        for text, width in texts_and_widths:
            chars = max(int(width * 1.15), 1)
            n = sum(max(1, -(-len(part) // chars)) for part in str(text or "").split("\n"))
            lines = max(lines, n)
        self.ws.row_dimensions[self.row].height = max(base, 13 * lines + 4)

    def width(self, first: int, last: int) -> float:
        return sum(self.ws.column_dimensions[get_column_letter(i)].width for i in range(first, last + 1))

    def title(self, text: str):
        self.row += 1
        self.cell(1, text, bold=True, fill=NAVY, color="FFFFFF", size=11, to=LAST,
                  align=Alignment(vertical="center"))
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1

    def headers(self, spans: list):
        for first, last, text in spans:
            self.cell(first, text, bold=True, fill=LIGHT, align=CENTER, to=last)
        self.height(*[(t, self.width(f, l)) for f, l, t in spans], base=30)
        self.row += 1


def build(db: Session, user: models.User, kpi: dict, prepared_by: models.User) -> bytes:
    """The report of `user` from their KPI detail (routers/kpi.py score(..., with_detail=True))."""
    today = date.today()
    manager = db.get(models.User, user.reports_to_id) if user.reports_to_id else None
    task_ids = [t["id"] for t in kpi["tasks"]]
    tasks = {t.id: t for t in db.query(models.Task).filter(models.Task.id.in_(task_ids)).all()} if task_ids else {}
    managed = {p["id"]: p for p in kpi["projects"]}
    project_ids = set(managed) | {t["project_id"] for t in kpi["tasks"] if t["project_id"]}
    projects = {p.id: p for p in db.query(models.Project).filter(models.Project.id.in_(project_ids)).all()} \
        if project_ids else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "KPI Report"
    widths = [5, 18, 28, 30, 30, 15, 10, 12, 12, 12, 11, 11, 9, 13, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    s = _Sheet(ws)

    # --- header, as the OKR sheet ---
    try:
        from openpyxl.drawing.image import Image  # needs Pillow; the report is complete without the logo
        logo = Image(str(LOGO))
        logo.height, logo.width = 52, int(52 * logo.width / logo.height)
        ws.add_image(logo, "A1")
    except Exception:
        pass
    s.cell(1, COMPANY, bold=True, size=16, align=CENTER, border=False, to=LAST, color=NAVY)
    ws.row_dimensions[1].height = 24
    s.row = 2
    s.cell(1, "TASK & PROJECT MANAGEMENT SYSTEM", bold=True, size=11, align=CENTER, border=False, to=LAST)
    s.row = 3
    s.cell(1, f"EMPLOYEE KPI REPORT — as of {_day(today)}", bold=True, size=12, align=CENTER, border=False, to=LAST)
    s.row = 4

    # --- the employee and their reporting manager ---
    s.title("EMPLOYEE INFORMATION")
    sbus = ", ".join(kpi["companies"]) or "—"
    manager_text = "—"
    if manager:
        manager_text = manager.name + "".join(f" · {x}" for x in (manager.employee_id, manager.designation) if x)
    pairs = [
        ("Employee Name", user.name), ("Employee ID", user.employee_id), ("Designation", user.designation or "—"),
        ("Role", _label(user.role)), ("SBU", sbus), ("Function", kpi["function"] or "—"),
        ("Department", kpi["department"] or "—"), ("Email", user.email), ("Reporting Manager", manager_text),
        ("Report Date", _day(today)), ("Period", f"Up to {_day(today)} (live KPI)"), ("Prepared By", prepared_by.name),
    ]
    blocks = [(1, 2, 3, 4), (5, 5, 6, 9), (10, 11, 12, LAST)]  # label from..to, value from..to
    for i in range(0, len(pairs), 3):
        texts = []
        for (lf, lt, vf, vt), (lab, val) in zip(blocks, pairs[i:i + 3]):
            s.cell(lf, lab, bold=True, fill=PALE, to=lt, align=Alignment(vertical="center", wrap_text=True))
            s.cell(vf, val, to=vt, align=Alignment(vertical="center", wrap_text=True))
            texts.append((val, s.width(vf, vt)))
        s.height(*texts, base=18)
        s.row += 1

    # --- the score ---
    s.title("1. KPI SUMMARY")
    s.headers([(1, 3, "Measure"), (4, 4, "Result"), (5, LAST, "How it is worked out")])
    rating = kpi["rating"]
    summary = [
        ("Total KPI", _num(kpi["total_kpi"]),
         "The Task KPI, the Project KPI, or their average when the employee has both. Target: 85 (Excellent)."),
        ("Rating", RATING_TEXT.get(rating, rating),
         "85 and above: Excellent · 70 – 84: Good · 50 – 69: Fair · below 50: Needs attention."),
        ("Task KPI", _num(kpi["task_kpi"]),
         "Delivery × 60% + On-Time Completion × 40%, over all the employee's tasks (each task once)."),
        ("Project KPI", _num(kpi["project_kpi"]), "The same formula over the projects the employee manages."),
        ("Delivery (Completion)", _num(kpi["delivery"]),
         "Average score of the counted tasks: a completed task scores 100; an open task scores 100 minus "
         "how far it is behind the progress expected by today."),
        ("On-Time Completion Rate", _pct(kpi["on_time"]),
         "Tasks completed on or before their due date ÷ (completed tasks + open tasks past their due date)."),
        ("On-Time Start Rate", _pct(kpi["otsr"]),
         "Tasks started on or before their planned start. For information - not part of the score."),
        ("Tasks Counted", kpi["assigned"], "Tasks whose planned start has passed, or that are completed."),
        ("Completed", kpi["completed"], f"Of which completed on time: {kpi['completed_on_time']}."),
        ("Past Due (open)", kpi["overdue"], "Open tasks whose due date has passed."),
        ("Not Counted Yet", kpi["not_counted"], "Planned start still ahead, or no due date: nothing to be behind yet."),
    ]
    for name, value, how in summary:
        s.cell(1, name, bold=True, to=3)
        fill = RATING_FILL.get(rating) if name in ("Total KPI", "Rating") else None
        s.cell(4, value, bold=name in ("Total KPI", "Rating"), align=CENTER, fill=fill)
        s.cell(5, how, to=LAST)
        s.height((how, s.width(5, LAST)))
        s.row += 1

    s.title("2. SCORE BY ROLE")
    s.headers([(1, 3, "Role"), (4, 4, "Items"), (5, 5, "Completed"), (6, 6, "Completed On Time"),
               (7, 7, "Past Due"), (8, 8, "Running"), (9, 9, "Not Counted"), (10, 10, "Delivery"),
               (11, 11, "On-Time %"), (12, 12, "KPI"), (13, LAST, "Rating")])
    for r in kpi["roles"]:
        unit = "projects" if r["unit"] == "projects" else "tasks"
        values = [(1, 3, f"{ROLE_TEXT.get(r['role'], r['role'])} ({unit})"), (4, 4, r["assigned"]),
                  (5, 5, r["completed"]), (6, 6, r["completed_on_time"]), (7, 7, r["overdue"]),
                  (8, 8, r["running"]), (9, 9, r["not_counted"]), (10, 10, _num(r["delivery"])),
                  (11, 11, _pct(r["on_time"])), (12, 12, _num(r["kpi"])),
                  (13, LAST, RATING_TEXT["no_tasks"] if r["kpi"] is None else
                   RATING_TEXT[_rating(r["kpi"])])]
        for first, last, v in values:
            s.cell(first, v, to=last, align=WRAP if first == 1 else CENTER, bold=first == 12)
        s.row += 1

    # --- projects: objective and expected outcome ---
    s.title("3. PROJECTS — OBJECTIVE, EXPECTED OUTCOME AND DELIVERY")
    s.headers([(1, 1, "SL"), (2, 2, "Project"), (3, 3, "Objective"), (4, 4, "Expected Outcome"),
               (5, 5, "Employee's Role"), (6, 6, "Priority"), (7, 7, "Employee's Tasks"), (8, 8, "Completed"),
               (9, 9, "Completed On Time"), (10, 10, "Project Completion"), (11, 11, "Due Date"),
               (12, 12, "Status"), (13, 13, "Score"), (14, LAST, "On-Time Completion Rate")])
    by_project: dict = {}
    for t in kpi["tasks"]:
        if t["project_id"]:
            by_project.setdefault(t["project_id"], []).append(t)
    order = sorted(project_ids, key=lambda pid: (projects[pid].code if pid in projects else "", pid))
    if not order:
        s.cell(1, "No projects.", to=LAST)
        s.row += 1
    for n, pid in enumerate(order, 1):
        p = projects.get(pid)
        if not p:
            continue
        mine = by_project.get(pid, [])
        roles = sorted({ROLE_TEXT[r] for t in mine for r in t["roles"]})
        if pid in managed:
            roles.insert(0, "Project Manager")
        counted = [t for t in mine if t["achievement"] is not None]
        if pid in managed:
            score = managed[pid]["achievement"]
        else:
            score = round(sum(t["achievement"] for t in counted) / len(counted), 1) if counted else None
        done = [t for t in mine if t["state"] == "completed"]
        good = sum(1 for t in done if t["on_time"] is not False)
        judged = len(done) + sum(1 for t in mine if t["state"] == "overdue")
        objective, outcome = _plain(p.objective) or "—", _plain(p.expected_outcome) or "—"
        name = f"{p.code}\n{p.name}"
        cells = [(1, 1, n), (2, 2, name), (3, 3, objective), (4, 4, outcome), (5, 5, ", ".join(roles) or "—"),
                 (6, 6, _label(p.priority)), (7, 7, len(mine)), (8, 8, len(done)), (9, 9, good),
                 (10, 10, _pct(round(p.completion_pct or 0, 1))),
                 (11, 11, _day(p.approved_due_date or p.baseline_due_date)), (12, 12, _label(p.status)),
                 (13, 13, _num(score)), (14, LAST, _pct(round(good / judged * 100, 1)) if judged else "—")]
        for first, last, v in cells:
            s.cell(first, v, to=last, align=WRAP if first in (2, 3, 4, 5) else CENTER, bold=first == 13)
        s.height((name, 18), (objective, 30), (outcome, 30), (cells[4][2], 15))
        s.row += 1

    # --- tasks ---
    s.title("4. TASKS — OBJECTIVE, EXPECTED OUTCOME AND ON-TIME COMPLETION")
    s.headers([(1, 1, "SL"), (2, 2, "Project"), (3, 3, "Task"), (4, 4, "Objective (Description)"),
               (5, 5, "Expected Outcome (Deliverable)"), (6, 6, "Employee's Role"), (7, 7, "Priority"),
               (8, 8, "Planned Start"), (9, 9, "Due Date"), (10, 10, "Completed On"), (11, 11, "Progress"),
               (12, 12, "Expected by Today"), (13, 13, "Score"), (14, 14, "On-Time Completion"),
               (15, 15, "Status")])
    rows = sorted(kpi["tasks"], key=lambda t: (t["project_code"] or "~", t["due_date"] or date.max, t["id"]))
    if not rows:
        s.cell(1, "No tasks.", to=LAST)
        s.row += 1
    for n, t in enumerate(rows, 1):
        task = tasks.get(t["id"])
        objective = _plain(task.description) if task else ""
        outcome = _plain(task.expected_deliverable) if task else ""
        project = f"{t['project_code']}\n{t['project_name']}" if t["project_code"] else "—"
        title = f"{t['code']}\n{t['title']}"
        verdict = _on_time(t)
        fill = {"On time": "C6EFCE", "Late": "FFC7CE", "Overdue": "FFC7CE"}.get(verdict)
        cells = [(1, n), (2, project), (3, title), (4, objective or "—"), (5, outcome or "—"),
                 (6, ", ".join(ROLE_TEXT[r] for r in t["roles"])), (7, _label(task.priority) if task else "—"),
                 (8, _day(t["planned_start_date"])), (9, _day(t["due_date"])), (10, _day(t["actual_due_date"])),
                 (11, _pct(t["progress_pct"])), (12, _pct(t["expected_pct"])), (13, _num(t["achievement"])),
                 (14, verdict), (15, _label(t["status"]))]
        for col, v in cells:
            s.cell(col, v, align=WRAP if col in (2, 3, 4, 5, 6) else CENTER, bold=col == 13,
                   fill=fill if col == 14 else None)
        s.height((project, 18), (title, 28), (objective, 30), (outcome, 30))
        s.row += 1

    # --- remarks and signatures ---
    s.title("5. REMARKS FROM REPORTING MANAGER")
    for _ in range(4):
        s.cell(1, None, to=LAST)
        ws.row_dimensions[s.row].height = 20
        s.row += 1

    s.row += 3
    sign = [(1, 4, "Employee Signature", user.name, user.designation),
            (11, LAST, "Reporting Manager Signature",
             manager.name if manager else "", manager.designation if manager else None)]
    top = Border(top=Side(style="thin", color="000000"))
    for first, last, caption, name, designation in sign:
        s.cell(first, caption, bold=True, border=False, to=last, align=Alignment(horizontal="center"))
        for i in range(first, last + 1):
            ws.cell(s.row, i).border = top
    s.row += 1
    for first, last, caption, name, designation in sign:
        s.cell(first, f"Name: {name or '______________________'}", border=False, to=last,
               align=Alignment(horizontal="center"))
    s.row += 1
    for first, last, caption, name, designation in sign:
        s.cell(first, f"Designation: {designation or '______________________'}", border=False, to=last,
               align=Alignment(horizontal="center"))
    s.row += 1
    for first, last, *_ in sign:
        s.cell(first, "Date: ____ / ____ / ________", border=False, to=last, align=Alignment(horizontal="center"))

    # --- printing: A4 landscape, all columns on one page width ---
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = f"KPI Report · {user.name} ({user.employee_id}) · {_day(today)}"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.sheet_view.showGridLines = False

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _rating(score: float) -> str:
    return "excellent" if score >= 85 else "good" if score >= 70 else "fair" if score >= 50 else "needs_attention"


def filename(user: models.User) -> str:
    safe = re.sub(r"[^A-Za-z0-9]+", "_", user.name).strip("_") or "employee"
    emp = re.sub(r"[^A-Za-z0-9-]+", "", user.employee_id or "")
    return f"KPI_Report_{emp}_{safe}_{date.today().isoformat()}.xlsx"
